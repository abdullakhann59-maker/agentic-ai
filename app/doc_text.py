"""Read text from documents (PDF, DOCX, TXT, MD, CSV). Shared by /knowledge and the document tools."""
from __future__ import annotations

from pathlib import Path


def read_document(path: str | Path) -> list[tuple[int, str]]:
    """Return a list of (page_number, text). Non-paged formats return one page."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".pdf":
        import pymupdf

        pages = []
        with pymupdf.open(p) as pdf:
            for i, page in enumerate(pdf, start=1):
                pages.append((i, page.get_text("text")))
        return pages
    if suffix == ".docx":
        import docx

        d = docx.Document(str(p))
        parts = [para.text for para in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return [(1, "\n".join(parts))]
    if suffix in (".txt", ".md", ".csv", ".json", ".html"):
        return [(1, p.read_text(encoding="utf-8", errors="ignore"))]
    raise ValueError(f"Unsupported file type: {suffix}. Use PDF, DOCX, TXT, MD, CSV, JSON or HTML.")


def pdf_has_text(path: str | Path) -> bool:
    return sum(len(t.strip()) for _, t in read_document(path)) > 50
